import docx
from fpdf import FPDF
import os

def create_docx(filename, content):
    doc = docx.Document()
    for line in content.split('\n'):
        doc.add_paragraph(line)
    doc.save(filename)
    print(f"Created {filename}")

def create_pdf(filename, content):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=11)
    for line in content.split('\n'):
        # Just simple mapping for FPDF, ignore special chars
        clean_line = line.encode('latin-1', 'replace').decode('latin-1')
        pdf.cell(200, 5, txt=clean_line, ln=1)
    pdf.output(filename)
    print(f"Created {filename}")

def generate_test_resumes():
    os.makedirs("test_resumes", exist_ok=True)
    
    # Candidate 1: Perfect Match (Senior Python Dev)
    cand1 = """Alice Smith
alicesmith@email.com | 555-123-4567 | San Francisco
    
EXPERIENCE
Senior Software Engineer
7 years of experience building scalable applications.
Skills: Python, Django, FastAPI, AWS, Docker, Machine Learning, SQL, Kubernetes, Linux.

EDUCATION
Master's in Computer Science
    
LOCATION
San Francisco, CA
"""

    # Candidate 2: Partial Match (Junior Python Dev)
    cand2 = """Bob Jones
bjones22@email.com | (555) 987-6543 | Remote

EXPERIENCE
Python Developer
3 years of experience.
Working heavily with Python, Flask, Pandas, Git, and HTML.

EDUCATION
Bachelor's Degree in IT
"""

    # Candidate 3: Poor Match (Frontend Dev)
    cand3 = """Charlie Brown
charlie.b@email.com | +1-555-666-7777

EXPERIENCE
Frontend Developer
5 years of experience in UI development.
Technologies: JavaScript, React, CSS, Node.js, TypeScript.

EDUCATION
B.S. in Software Engineering
Location: Remote
"""

    create_docx("test_resumes/Alice_Smith_Resume.docx", cand1)
    create_pdf("test_resumes/Bob_Jones_Resume.pdf", cand2)
    create_docx("test_resumes/Charlie_Brown_Frontend.docx", cand3)

if __name__ == "__main__":
    generate_test_resumes()
    print("Test resumes generated in 'test_resumes' folder.")
